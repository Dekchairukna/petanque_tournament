"""
นำเข้ารายชื่อทีม/ประเภทการแข่งขันจากไฟล์ Excel และส่งออกตารางแบ่งสายเป็น Excel

รูปแบบไฟล์นำเข้าที่รองรับ
1) แบบบล็อก (เหมือนไฟล์รายชื่อ อปท.)
       ประเภทบุคคลชาย รุ่นอายุ 12 ปี
       ที่ | ชื่ออปท. | อำเภอ | จังหวัด
       1  | เทศบาล...| ...   | ...
       ...
       ประเภทบุคคลหญิง รุ่นอายุ 12 ปี
       ...
2) แบบตาราง (1 แถว = 1 ทีม) มีหัวคอลัมน์ "ประเภท" และ "ชื่อ..." / "ทีม" / "หน่วยงาน"
   (จะมีคอลัมน์ "รุ่น" / "อายุ", "อำเภอ", "จังหวัด" ด้วยหรือไม่ก็ได้)
"""
import io
import os
import re
from collections import Counter, OrderedDict

THAI_DIGITS = str.maketrans("๐๑๒๓๔๕๖๗๘๙", "0123456789")


# ------------------------------------------------------------------ reading
def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def read_workbook_rows(file_bytes, filename):
    """คืนค่า list ของ (ชื่อชีต, rows) โดย rows เป็น list ของ list ข้อความ"""
    name = (filename or "").lower()
    sheets = []
    if name.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=file_bytes)
        for sh in book.sheets():
            rows = []
            for r in range(sh.nrows):
                rows.append([_cell_text(sh.cell_value(r, c)) for c in range(sh.ncols)])
            sheets.append((sh.name, rows))
    elif name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        for ws in wb.worksheets:
            rows = [[_cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
            sheets.append((ws.title, rows))
        wb.close()
    elif name.endswith(".csv"):
        import csv

        text = file_bytes.decode("utf-8-sig", errors="replace")
        rows = [[_cell_text(v) for v in row] for row in csv.reader(io.StringIO(text))]
        sheets.append(("csv", rows))
    else:
        raise ValueError("รองรับเฉพาะไฟล์ .xls, .xlsx หรือ .csv")
    return sheets


# ------------------------------------------------------------------ helpers
ORG_ABBREVIATIONS = [
    ("องค์การบริหารส่วนจังหวัด", "อบจ."),
    ("องค์การบริหารส่วนตำบล", "อบต."),
    ("เทศบาลนคร", "ทน."),
    ("เทศบาลเมือง", "ทม."),
    ("เทศบาลตำบล", "ทต."),
]


def abbreviate_org(name):
    text = (name or "").strip()
    for full, short in ORG_ABBREVIATIONS:
        if text.startswith(full):
            return short + text[len(full):].strip()
    return text


CATEGORY_ORDER = ["เดี่ยว", "คู่", "คู่ผสม", "ทีม", "ชู้ตติ้ง", "อื่นๆ"]

# ค่าจากตารางสรุปการกำหนดจำนวนสาย “อปท.สกลนคร 8-18 ต.ค.69”
# ใช้เป็นค่าสำรองเมื่อไฟล์รายชื่อไม่มีชีตสรุปจำนวนสายอยู่ภายใน
DEFAULT_GROUP_COUNTS = {
    ("เดี่ยว", 12, "ชาย"): 16, ("เดี่ยว", 12, "หญิง"): 12,
    ("เดี่ยว", 14, "ชาย"): 12, ("เดี่ยว", 14, "หญิง"): 8,
    ("เดี่ยว", 16, "ชาย"): 8,  ("เดี่ยว", 16, "หญิง"): 8,
    ("เดี่ยว", 18, "ชาย"): 8,  ("เดี่ยว", 18, "หญิง"): 6,
    ("คู่", 12, "ชาย"): 16,    ("คู่", 12, "หญิง"): 12,
    ("คู่", 14, "ชาย"): 12,    ("คู่", 14, "หญิง"): 8,
    ("คู่", 16, "ชาย"): 8,     ("คู่", 16, "หญิง"): 8,
    ("คู่", 18, "ชาย"): 8,     ("คู่", 18, "หญิง"): 6,
    ("คู่ผสม", 12, "ผสม"): 12, ("คู่ผสม", 14, "ผสม"): 8,
    ("คู่ผสม", 16, "ผสม"): 6,  ("คู่ผสม", 18, "ผสม"): 6,
    ("ทีม", 12, "ชาย"): 16,   ("ทีม", 12, "หญิง"): 12,
    ("ทีม", 14, "ชาย"): 10,   ("ทีม", 14, "หญิง"): 8,
    ("ทีม", 16, "ชาย"): 8,    ("ทีม", 16, "หญิง"): 6,
    ("ทีม", 18, "ชาย"): 6,    ("ทีม", 18, "หญิง"): 6,
}


def default_group_count(event):
    return DEFAULT_GROUP_COUNTS.get((event.get("category"), event.get("age"), event.get("gender")))


def classify_event(title):
    """แยก ประเภท / เพศ / อายุ ออกจากชื่อประเภทการแข่งขัน"""
    t = (title or "").translate(THAI_DIGITS)
    if "ชู้ต" in t or "ยิงเป้า" in t or "shoot" in t.lower():
        category = "ชู้ตติ้ง"
    elif "ผสม" in t:
        category = "คู่ผสม"
    elif "ทีม" in t:
        category = "ทีม"
    elif "คู่" in t:
        category = "คู่"
    elif "บุคคล" in t or "เดี่ยว" in t:
        category = "เดี่ยว"
    else:
        category = "อื่นๆ"

    if "ผสม" in t:
        gender = "ผสม"
    elif "หญิง" in t:
        gender = "หญิง"
    elif "ชาย" in t:
        gender = "ชาย"
    else:
        gender = "-"

    m = re.search(r"(\d+)\s*ปี", t)
    age = int(m.group(1)) if m else None
    return category, gender, age


def clean_event_title(raw):
    t = re.sub(r"\s+", " ", (raw or "").translate(THAI_DIGITS)).strip()
    t = re.sub(r"^ประเภท\s*", "", t)
    # ไฟล์บางชุดใส่จำนวนทีมต่อท้ายหัวข้อ เช่น "/ 54 ทีม" หรือ "/ 28"
    # จำนวนนี้เป็นข้อมูลประกอบ ไม่ใช่ส่วนหนึ่งของชื่อประเภทการแข่งขัน
    t = re.sub(r"\s*/\s*\d+\s*(?:ทีม)?\s*$", "", t).strip()
    if t.startswith("ผสม"):
        t = "คู่" + t
    return t.strip() or raw


def export_event_label(title):
    """ข้อความแสดงบนหัวสาย เช่น 'ประเภท บุคคลชาย รุ่นอายุไม่เกิน 12 ปี'"""
    t = clean_event_title(title)
    t = re.sub(r"รุ่นอายุ\s*(?!ไม่เกิน)(\d+)", r"รุ่นอายุไม่เกิน \1", t)
    return f"ประเภท {t}"


def _is_event_header(cells):
    non_empty = [c for c in cells if c]
    if not non_empty:
        return False
    first = cells[0] if cells else ""
    # อนุญาตให้แถวหัวข้อมีจำนวนทีม/จำนวนสายอยู่ในคอลัมน์ถัดไปด้วย
    return first.startswith("ประเภท") and len(first) > len("ประเภท") + 1


def _first_int(value):
    m = re.search(r"\d+", (value or "").translate(THAI_DIGITS))
    return int(m.group()) if m else None


def _summary_group_counts(sheets):
    """อ่านตารางสรุป ประเภท/อายุ/เพศ/จำนวนทีม/จำนวนสาย จากทุกชีต"""
    hints = {}
    for _sheet_name, rows in sheets:
        for i, cells in enumerate(rows):
            if not any(c == "ประเภท" for c in cells):
                continue
            category_col = next((j for j, c in enumerate(cells) if c == "ประเภท"), None)
            age_col = next((j for j, c in enumerate(cells) if "รุ่นอายุ" in c), None)
            gender_col = next((j for j, c in enumerate(cells) if c == "เพศ"), None)
            team_col = next((j for j, c in enumerate(cells) if c == "จำนวนทีม"), None)
            if None in (category_col, age_col, gender_col, team_col):
                continue

            group_col = next((j for j, c in enumerate(cells) if "จำนวนสาย" in c), None)
            data_start = i + 1
            if group_col is None and i + 1 < len(rows):
                group_col = next((j for j, c in enumerate(rows[i + 1]) if "จำนวนสาย" in c), None)
                if group_col is not None:
                    data_start = i + 2
            if group_col is None:
                continue

            current_category = ""
            for row in rows[data_start:]:
                if any(c == "ประเภท" for c in row):
                    break
                raw_category = _get(row, category_col)
                if raw_category:
                    current_category = raw_category
                age = _first_int(_get(row, age_col))
                gender = _get(row, gender_col)
                team_count = _first_int(_get(row, team_col))
                group_count = _first_int(_get(row, group_col))
                if not current_category or not age or not gender or not team_count or not group_count:
                    continue
                category, _, _ = classify_event(current_category)
                hints.setdefault((category, age, gender), []).append((team_count, group_count))
    return hints


def _find_col(header, *keywords, exclude=()):
    for idx, h in enumerate(header):
        if any(k in h for k in keywords) and not any(x in h for x in exclude):
            return idx
    return None


def _header_columns(cells):
    """ถ้าแถวนี้เป็นหัวตาราง คืน dict ตำแหน่งคอลัมน์"""
    exact = {"ทีม", "หน่วยงาน", "สังกัด", "ชื่อ", "Team", "team", "Name", "name"}
    name_col = next(
        (i for i, c in enumerate(cells) if c and (c.startswith("ชื่อ") or c in exact)),
        None,
    )
    if name_col is None:
        return None
    seq_col = _find_col(cells, "ลำดับ")
    if seq_col is None:
        seq_col = next((i for i, c in enumerate(cells) if c in ("ที่", "No", "no", "#")), None)
    return {
        "seq": seq_col,
        "name": name_col,
        "district": _find_col(cells, "อำเภอ", "เขต"),
        "province": _find_col(cells, "จังหวัด"),
        "event": next((i for i, c in enumerate(cells) if c.startswith("ประเภท") and len(c) <= 14), None),
        "age": _find_col(cells, "รุ่น", "อายุ"),
    }


def _get(row, idx):
    if idx is None or idx >= len(row):
        return ""
    return row[idx]


def parse_events(sheets):
    """
    คืนค่า dict:
      title_lines: บรรทัดชื่อรายการ (จากชีตที่ไม่มีรายชื่อทีม)
      events: [{key, raw_title, title, category, gender, age, teams:[{name, district, province}]}]
    """
    events = OrderedDict()
    title_lines = []
    group_count_hints = _summary_group_counts(sheets)

    def get_event(raw_title):
        title = clean_event_title(raw_title)
        key = title
        if key not in events:
            category, gender, age = classify_event(title)
            events[key] = {
                "key": key,
                "raw_title": raw_title,
                "title": title,
                "category": category,
                "gender": gender,
                "age": age,
                "teams": [],
            }
        return events[key]

    for sheet_name, rows in sheets:
        current = None
        cols = None
        flat = False
        sheet_had_event = False

        for cells in rows:
            if not any(cells):
                continue

            if _is_event_header(cells):
                current = get_event(next(c for c in cells if c))
                row_group_count = next((_first_int(c) for c in cells[1:] if re.search(r"\d+\s*สาย", c or "")), None)
                if row_group_count:
                    current["suggested_group_count"] = row_group_count
                    current["group_count_source"] = "file"
                cols = None
                flat = False
                sheet_had_event = True
                continue

            header = _header_columns(cells)
            if header and not any(re.fullmatch(r"\d+", c or "") for c in cells[:1]):
                # หัวตาราง
                cols = header
                flat = header["event"] is not None
                continue

            if flat and cols:
                ev_text = _get(cells, cols["event"])
                name = _get(cells, cols["name"])
                if not ev_text or not name:
                    continue
                age_text = _get(cells, cols["age"]) if cols["age"] not in (None, cols["event"]) else ""
                if age_text and age_text not in ev_text:
                    if re.fullmatch(r"\d+", age_text.translate(THAI_DIGITS)):
                        age_text = f"รุ่นอายุ {age_text} ปี"
                    ev_text = f"{ev_text} {age_text}"
                ev = get_event(ev_text)
                sheet_had_event = True
                ev["teams"].append({
                    "name": name,
                    "district": _get(cells, cols["district"]),
                    "province": _get(cells, cols["province"]),
                })
                continue

            if current is not None:
                if cols:
                    name = _get(cells, cols["name"])
                    seq = _get(cells, cols["seq"]) if cols["seq"] is not None else "1"
                    if not name or not seq:
                        continue
                    current["teams"].append({
                        "name": name,
                        "district": _get(cells, cols["district"]),
                        "province": _get(cells, cols["province"]),
                    })
                else:
                    # ไม่มีหัวตาราง: ถ้าคอลัมน์แรกเป็นตัวเลข ใช้คอลัมน์ถัดไปเป็นชื่อ
                    first = cells[0].translate(THAI_DIGITS) if cells else ""
                    if re.fullmatch(r"\d+", first) and len(cells) > 1 and cells[1]:
                        current["teams"].append({
                            "name": cells[1],
                            "district": _get(cells, 2),
                            "province": _get(cells, 3),
                        })
                    elif not re.fullmatch(r"\d+", first) and first:
                        current["teams"].append({"name": first, "district": "", "province": ""})
                continue

        if not sheet_had_event:
            for cells in rows:
                line = " ".join(c for c in cells if c).strip()
                if line:
                    title_lines.append(line)

    result = [ev for ev in events.values() if ev["teams"]]
    for ev in result:
        key = (ev["category"], ev["age"], ev["gender"])
        candidates = group_count_hints.get(key) or []
        if candidates:
            # ใช้จำนวนทีมที่ใกล้กับรายชื่อจริงที่สุด เผื่อหัวตารางยังเป็นยอดเดิม
            # เช่น ระบุ 54 ทีม แต่มีการเพิ่มรายชื่อภายหลังเป็น 55 ทีม
            _, group_count = min(candidates, key=lambda item: abs(item[0] - len(ev["teams"])))
            ev["suggested_group_count"] = group_count
            ev["group_count_source"] = "file"
    return {"title_lines": title_lines, "events": result}


def event_sort_key(ev):
    cat = ev.get("category") or "อื่นๆ"
    return (
        CATEGORY_ORDER.index(cat) if cat in CATEGORY_ORDER else 99,
        ev.get("age") or 999,
        {"ชาย": 0, "หญิง": 1}.get(ev.get("gender"), 2),
    )


def prepare_team_names(teams, abbreviate=True):
    """คืน list ของ dict ที่มี display (ชื่อที่ใช้ในระบบ ไม่ซ้ำกัน)"""
    prepared = []
    for t in teams:
        display = abbreviate_org(t["name"]) if abbreviate else t["name"].strip()
        prepared.append({**t, "display": display})
    counts = Counter(p["display"] for p in prepared)
    seen = Counter()
    for p in prepared:
        if counts[p["display"]] > 1:
            seen[p["display"]] += 1
            p["display"] = f"{p['display']} {seen[p['display']]}"
    return prepared


# ------------------------------------------------------------------ export
def safe_sheet_title(title, used):
    t = re.sub(r"[\[\]\:\*\?\/\\]", " ", title)
    t = t.replace("รุ่นอายุ", "").replace("ประเภท", "")
    t = re.sub(r"\s+", " ", t).strip()[:31] or "Sheet"
    base, n = t, 2
    while t in used:
        suffix = f" ({n})"
        t = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(t)
    return t


def build_draw_workbook(events, subtitle, left_logo=None, right_logo=None, main_title="ตารางแบ่งสายการแข่งขัน"):
    """
    events: [{
        title, label,
        rounds: [{label, round_type, groups: [{group_no, slots: [{no, name, court, scores:[..]}], qualified:[...]}]}],
        teams: [{no, display, name, district, province, group_no}]
    }]
    คืน BytesIO ของไฟล์ .xlsx ตามรูปแบบตารางแบ่งสาย
    """
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, Side
    from openpyxl.worksheet.pagebreak import Break

    thin = Side(style="thin")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    FONT = "Calibri"
    f_title = Font(name=FONT, size=22, bold=True)
    f_sub = Font(name=FONT, size=20, bold=True)
    f_bold = Font(name=FONT, size=16, bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)

    wb = Workbook()
    wb.remove(wb.active)
    used_titles = set()

    logo_cache = {}

    def shrink(src, max_h):
        """ย่อรูปให้เล็กพอดีกับหัวกระดาษ (กันไฟล์ใหญ่เพราะโลโก้ซ้ำทุกหน้า)"""
        key = (id(src), max_h)
        if key in logo_cache:
            return logo_cache[key]
        from PIL import Image as PILImage

        if isinstance(src, (bytes, bytearray, memoryview)):
            pil = PILImage.open(io.BytesIO(bytes(src)))
        elif isinstance(src, str) and os.path.exists(src):
            pil = PILImage.open(src)
        else:
            logo_cache[key] = None
            return None
        pil = pil.convert("RGBA")
        target_h = max_h * 2  # 2 เท่าให้พิมพ์คมชัด
        if pil.height > target_h:
            pil = pil.resize((max(1, int(pil.width * target_h / pil.height)), target_h), PILImage.LANCZOS)
        buf = io.BytesIO()
        pil.save(buf, format="PNG", optimize=True)
        logo_cache[key] = (buf.getvalue(), pil.width, pil.height)
        return logo_cache[key]

    def add_logo(ws, src, anchor, max_h):
        # src = bytes ของรูป หรือ path ไฟล์
        if not src:
            return
        try:
            shrunk = shrink(src, max_h)
            if not shrunk:
                return
            data, w, h = shrunk
            img = XLImage(io.BytesIO(data))
        except Exception:
            return
        img.height = max_h
        img.width = int(w * max_h / float(h))
        ws.add_image(img, anchor)

    def boxed(ws, r1, c1, r2, c2, value=None, font=f_bold, align=center):
        if (r1, c1) != (r2, c2):
            ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
        for rr in range(r1, r2 + 1):
            for cc in range(c1, c2 + 1):
                cell = ws.cell(row=rr, column=cc)
                cell.border = border
                cell.font = font
                cell.alignment = align
        ws.cell(row=r1, column=c1, value=value)

    def write_group(ws, r, c0, event_label, round_label, round_type, group):
        # แถวหัวสาย
        # ชื่อประเภท + รอบ รวมใน A:C (ชื่อประเภทยาว เช่น ทีม 3 คนหญิง จะได้ไม่ถูกตัด)
        ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
        cell = ws.cell(row=r, column=c0, value=f"{event_label}    {round_label}")
        cell.font = f_bold
        cell.alignment = Alignment(horizontal="left", vertical="center", shrink_to_fit=True)
        cell = ws.cell(row=r, column=c0 + 3, value="สายที่")
        cell.font = f_bold
        cell.alignment = Alignment(horizontal="right", vertical="center")
        cell = ws.cell(row=r, column=c0 + 4, value=group["group_no"])
        cell.font = f_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")

        # หัวตาราง
        boxed(ws, r + 1, c0, r + 2, c0, "สนาม")
        boxed(ws, r + 1, c0 + 1, r + 2, c0 + 1, "ลำดับที่")
        boxed(ws, r + 1, c0 + 2, r + 2, c0 + 2, "รายชื่อทีม")
        double = round_type == "double_knockout"
        if double:
            boxed(ws, r + 1, c0 + 3, r + 1, c0 + 5, "แข่งขันครั้งที่")
            for i in range(3):
                boxed(ws, r + 2, c0 + 3 + i, r + 2, c0 + 3 + i, i + 1)
        else:
            boxed(ws, r + 1, c0 + 3, r + 2, c0 + 5, "ผลการแข่งขัน")

        # ทีม
        slots = group["slots"]
        rr = r + 3
        for idx, slot in enumerate(slots):
            if idx % 2 == 0:
                end = rr + (1 if idx + 1 < len(slots) else 0)
                boxed(ws, rr, c0, end, c0, slot.get("court") or None)
            boxed(ws, rr, c0 + 1, rr, c0 + 1, slot["no"])
            boxed(ws, rr, c0 + 2, rr, c0 + 2, slot["name"], align=left)
            scores = slot.get("scores") or []
            if double:
                for i in range(3):
                    val = scores[i] if i < len(scores) else None
                    boxed(ws, rr, c0 + 3 + i, rr, c0 + 3 + i, val)
            else:
                boxed(ws, rr, c0 + 3, rr, c0 + 5, scores[0] if scores else None)
            rr += 1

        q = group.get("qualified") or []
        if double:
            a = q[0] if len(q) > 0 else "….................................."
            b = q[1] if len(q) > 1 else ".............................................."
            text = f"ทีมที่เข้ารอบ 1. {a}      2. {b}"
        else:
            a = q[0] if q else "............................................................"
            text = f"ทีมผู้ชนะ {a}"
        ws.cell(row=rr, column=c0, value=text).font = f_bold
        return rr  # แถวสุดท้ายที่ใช้

    for ev in events:
        ws = wb.create_sheet(safe_sheet_title(ev["title"], used_titles))
        widths = {"A": 19.16, "B": 20.83, "C": 30.83, "D": 9, "E": 9, "F": 9, "G": 5,
                  "H": 19.16, "I": 20.83, "J": 30.83, "K": 10.83, "L": 9, "M": 9}
        for col, w in widths.items():
            ws.column_dimensions[col].width = w

        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = ws.page_margins.bottom = 0.5
        ws.print_options.horizontalCentered = True

        row = 1
        for rnd in ev["rounds"]:
            groups = rnd["groups"]
            if not groups:
                continue
            for p in range(0, len(groups), 4):
                page_groups = groups[p:p + 4]
                # หัวกระดาษ
                ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=13)
                c = ws.cell(row=row, column=1, value=main_title)
                c.font, c.alignment = f_title, center
                ws.row_dimensions[row].height = 35
                ws.merge_cells(start_row=row + 1, start_column=1, end_row=row + 1, end_column=13)
                c = ws.cell(row=row + 1, column=1, value=subtitle)
                c.font, c.alignment = f_sub, center
                ws.row_dimensions[row + 1].height = 40
                add_logo(ws, left_logo, f"A{row}", 95)
                add_logo(ws, right_logo, f"M{row}", 90)
                r = row + 2

                for pi in range(0, len(page_groups), 2):
                    if pi > 0:
                        ws.row_dimensions[r].height = 35
                        r += 1
                    pair = page_groups[pi:pi + 2]
                    last = r
                    for side, grp in enumerate(pair):
                        end = write_group(ws, r, 1 + side * 7, ev["label"], rnd["label"], rnd["round_type"], grp)
                        last = max(last, end)
                    for rr in range(r, last + 1):
                        ws.row_dimensions[rr].height = 35
                    r = last + 1

                ws.row_breaks.append(Break(id=r - 1))
                row = r

    # ชีตรายชื่อทีมทั้งหมด (อ้างอิงเลขลำดับ)
    ws = wb.create_sheet("รายชื่อทีม")
    headers = ["ประเภท", "สายที่", "ลำดับที่", "ชื่อทีม", "ชื่อหน่วยงาน (เต็ม)", "อำเภอ", "จังหวัด"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.alignment = center
        from openpyxl.styles import PatternFill

        c.fill = PatternFill("solid", fgColor="1F4E78")
    r = 2
    for ev in events:
        for t in ev.get("teams", []):
            ws.append([ev["title"], t.get("group_no"), t.get("no"), t.get("display"),
                       t.get("name"), t.get("district"), t.get("province")])
            r += 1
    for col, w in zip("ABCDEFG", [34, 8, 9, 30, 38, 18, 14]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    if r > 2:
        ws.auto_filter.ref = f"A1:G{r - 1}"

    out = io.BytesIO()
    wb.save(out)
    return dedupe_media(out.getvalue())


def dedupe_media(xlsx_bytes):
    """openpyxl เก็บรูปซ้ำทุกครั้งที่วาง -> รวมรูปที่เหมือนกันให้เหลือไฟล์เดียว (ไฟล์เล็กลงมาก)"""
    import hashlib
    import zipfile

    src = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    names = src.namelist()
    first_by_hash, rename = {}, {}
    for n in names:
        if n.startswith("xl/media/"):
            h = hashlib.sha1(src.read(n)).hexdigest()
            if h in first_by_hash:
                rename[n] = first_by_hash[h]
            else:
                first_by_hash[h] = n
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for n in names:
            if n in rename:
                continue
            data = src.read(n)
            if rename and n.startswith("xl/drawings/_rels/"):
                text = data.decode("utf-8")
                for old, new in rename.items():
                    text = text.replace("/" + old + '"', "/" + new + '"')
                data = text.encode("utf-8")
            dst.writestr(n, data)
    out.seek(0)
    return out
